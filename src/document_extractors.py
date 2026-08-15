from pathlib import Path

from docx import Document as WordDocument
from pptx import Presentation
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".pdf",
    ".docx",
    ".pptx",
}


def extract_text_from_txt(file_path: Path) -> str:
    """Extrae texto de archivos TXT o Markdown."""

    try:
        return file_path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return file_path.read_text(encoding="latin-1")


def extract_text_from_pdf(file_path: Path) -> str:
    """Extrae el texto digital de un PDF."""

    reader = PdfReader(str(file_path))
    pages: list[str] = []

    for page_number, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""

        if page_text.strip():
            pages.append(
                f"--- Página {page_number} ---\n{page_text.strip()}"
            )

    return "\n\n".join(pages)


def extract_text_from_docx(file_path: Path) -> str:
    """Extrae párrafos y tablas de un archivo Word."""

    document = WordDocument(str(file_path))
    content: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            style_name = str(getattr(paragraph.style, "name", "") or "")
            if style_name.casefold().startswith("heading"):
                level_match = "".join(character for character in style_name if character.isdigit())
                level = max(1, min(3, int(level_match or "1")))
                content.append(f"{'#' * level} {text}")
            else:
                content.append(text)

    for table_number, table in enumerate(document.tables, start=1):
        content.append(f"--- Tabla {table_number} ---")

        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            if any(cells):
                content.append(" | ".join(cells))

    return "\n".join(content)


def extract_text_from_pptx(file_path: Path) -> str:
    """Extrae textos y tablas de una presentación PowerPoint."""

    presentation = Presentation(str(file_path))
    content: list[str] = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_content: list[str] = []

        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()

                if text:
                    slide_content.append(f"# {text}" if shape == slide.shapes.title else text)

            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [
                        cell.text.strip()
                        for cell in row.cells
                    ]

                    if any(cells):
                        slide_content.append(" | ".join(cells))

        if slide_content:
            content.append(
                f"--- Diapositiva {slide_number} ---\n"
                + "\n".join(slide_content)
            )

    return "\n\n".join(content)


def extract_document_text(file_path: Path) -> str:
    """
    Detecta el formato del archivo y extrae su texto.

    Lanza ValueError cuando el formato no es compatible.
    """

    extension = file_path.suffix.lower()

    if extension in {".txt", ".md"}:
        return extract_text_from_txt(file_path)

    if extension == ".pdf":
        return extract_text_from_pdf(file_path)

    if extension == ".docx":
        return extract_text_from_docx(file_path)

    if extension == ".pptx":
        return extract_text_from_pptx(file_path)

    raise ValueError(
        f"Formato no compatible: {extension or 'sin extensión'}"
    )
